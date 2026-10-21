"""Text, DOCX, PDF and image parsers. PDFs without a text layer and images go through OCR."""

from __future__ import annotations

import io
import re
from typing import Any

from app.agent.parsers.base import ParsedDocument, Segment, build_table, markdown_table
from app.ai.ocr.base import OCRUnavailable
from app.core.services.files import decode_text

OCR_DPI = 300


def _detect_pasted_table(text: str) -> list[list[str]] | None:
    """Pasted spreadsheet ranges arrive as TSV; also accept markdown and semicolon tables."""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) < 2:
        return None
    for splitter in ("\t", ";"):
        counts = [ln.count(splitter) for ln in lines]
        if min(counts) >= 1 and max(counts) - min(counts) <= 1 and sum(c >= 1 for c in counts) >= 2:
            return [[c.strip() for c in ln.split(splitter)] for ln in lines]
    if all(ln.strip().startswith("|") for ln in lines):
        rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
        return [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
    return None


def parse_text(text: str, name: str = "text") -> ParsedDocument:
    table = _detect_pasted_table(text)
    if table is not None:
        built = build_table(table)
        if built is not None:
            headers, rows, info = built
            doc = ParsedDocument(parser="pasted_table", is_tabular=True)
            doc.add_table(name="Pasted table", headers=headers, rows=rows, extra=info)
            doc.meta["full_text"] = markdown_table(headers, [r for _, r in rows])
            return doc
    doc = ParsedDocument(parser="text")
    doc.segments.append(Segment("text", {"text": text}, page=1))
    doc.meta["full_text"] = text
    doc.meta["page_offsets"] = [[1, 0, len(text)]]
    return doc


def parse_text_file(data: bytes, filename: str) -> ParsedDocument:
    text, encoding = decode_text(data)
    doc = parse_text(text, filename)
    doc.meta["encoding"] = encoding
    return doc


def parse_docx(data: bytes, filename: str) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    doc = ParsedDocument(parser="python-docx")
    parts: list[str] = []
    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(block, document).text.strip()
            if text:
                parts.append(text)
        elif tag == "tbl":
            table = Table(block, document)
            # python-docx repeats a merged cell's value in every column it spans (fill-down behaviour).
            raw = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            built = build_table(raw)
            if built:
                headers, rows, info = built
                doc.add_table(name=f"Table {sum(1 for s in doc.segments if s.kind == 'table') + 1}", headers=headers,
                              rows=rows, extra=info)
                parts.append(markdown_table(headers, [r for _, r in rows]))
            else:
                parts.append("\n".join(" | ".join(r) for r in raw))
    text = "\n\n".join(parts)
    doc.segments.insert(0, Segment("text", {"text": text}, page=1))
    for s in doc.segments:  # keep table_seq pointing at the right segment after the insert
        if s.table_seq is not None:
            s.table_seq += 1
    doc.meta["full_text"] = text
    doc.meta["page_offsets"] = [[1, 0, len(text)]]
    return doc


def _ocr_image(image) -> dict[str, Any]:
    from app.ai.ocr.tesseract import get_ocr

    result = get_ocr().recognize(image)
    return {
        "text": result.text, "width": result.width, "height": result.height, "rotation": result.rotation,
        "lines": [{"text": ln.text, "bbox": [round(v, 1) for v in ln.bbox], "conf": round(ln.confidence, 1)} for ln in result.lines],
        "mean_confidence": result.mean_confidence, "engine": result.engine,
    }


def parse_pdf(data: bytes, filename: str) -> ParsedDocument:
    import pdfplumber
    import pypdfium2

    doc = ParsedDocument(parser="pdfplumber")
    parts: list[str] = []
    offsets: list[list[int]] = []
    confidences: list[float] = []
    cursor = 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        doc.page_count = len(pdf.pages)
        pdfium = None
        for number, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            content: dict[str, Any] = {"width": float(page.width), "height": float(page.height)}
            ocr_conf = None
            if len(text) >= 20:
                content["lines"] = [{"text": ln["text"], "bbox": [round(ln["x0"], 1), round(ln["top"], 1), round(ln["x1"], 1),
                                                                   round(ln["bottom"], 1)]} for ln in page.extract_text_lines()]
                for raw in page.extract_tables() or []:
                    built = build_table(raw)
                    if built:
                        headers, rows, info = built
                        doc.add_table(name=f"Page {number} table", headers=headers, rows=rows, page=number, extra=info)
                        text += "\n\n" + markdown_table(headers, [r for _, r in rows])
            else:
                try:
                    pdfium = pdfium or pypdfium2.PdfDocument(data)
                    image = pdfium[number - 1].render(scale=OCR_DPI / 72).to_pil()
                    ocr = _ocr_image(image)
                except OCRUnavailable as exc:
                    doc.warnings.append(f"Page {number} has no text layer and OCR is unavailable: {exc}")
                    continue
                scale = 72 / OCR_DPI
                text = ocr["text"]
                content["lines"] = [{**ln, "bbox": [round(v * scale, 1) for v in ln["bbox"]]} for ln in ocr["lines"]]
                content["ocr"] = True
                ocr_conf = ocr["mean_confidence"]
                doc.ocr_used = True
                if ocr_conf is not None:
                    confidences.append(ocr_conf)
                if not text.strip():
                    doc.warnings.append(f"Page {number}: OCR found no readable text")
            content["text"] = text
            page_text = (f"--- Page {number} ---\n" if doc.page_count > 1 else "") + text
            offsets.append([number, cursor, cursor + len(page_text)])
            parts.append(page_text)
            cursor += len(page_text) + 2
            doc.segments.insert(sum(1 for s in doc.segments if s.kind == "text"),
                                Segment("text", content, page=number, ocr_confidence=ocr_conf))
        if pdfium is not None:
            pdfium.close()
    # Recompute table_seq after text segments were inserted ahead of tables.
    _reindex(doc)
    doc.ocr_confidence = round(sum(confidences) / len(confidences), 1) if confidences else None
    if doc.ocr_used:
        doc.parser = "pdfplumber+tesseract"
    doc.meta["full_text"] = "\n\n".join(parts)
    doc.meta["page_offsets"] = offsets
    return doc


def _reindex(doc: ParsedDocument) -> None:
    current = None
    for i, s in enumerate(doc.segments):
        if s.kind == "table":
            current = i
        elif s.kind == "row":
            s.table_seq = current


def parse_image(data: bytes, filename: str) -> ParsedDocument:
    from PIL import Image

    image = Image.open(io.BytesIO(data))
    image.load()
    doc = ParsedDocument(parser="tesseract", ocr_used=True, page_count=1)
    try:
        ocr = _ocr_image(image)
    except OCRUnavailable as exc:
        doc.warnings.append(f"OCR is unavailable: {exc}")
        doc.meta["full_text"] = ""
        return doc
    doc.ocr_confidence = ocr["mean_confidence"]
    doc.segments.append(Segment("text", {"text": ocr["text"], "lines": ocr["lines"], "width": ocr["width"],
                                         "height": ocr["height"], "ocr": True, "rotation": ocr["rotation"]},
                                page=1, ocr_confidence=ocr["mean_confidence"]))
    if ocr["mean_confidence"] is not None and ocr["mean_confidence"] < 70:
        doc.warnings.append(f"Low OCR confidence ({ocr['mean_confidence']}%): check every extracted value against the image")
    if not ocr["text"].strip():
        doc.warnings.append("OCR found no readable text in the image")
    doc.meta["full_text"] = ocr["text"]
    doc.meta["page_offsets"] = [[1, 0, len(ocr["text"])]]
    return doc
