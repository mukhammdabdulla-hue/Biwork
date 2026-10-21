"""CSV / XLSX / XLS parsers: sheets, header-row detection, typed cells, merged cells."""

from __future__ import annotations

import csv
import io
from typing import Any

from app.agent.parsers.base import ParsedDocument, build_table, markdown_table
from app.core.services.files import decode_text

FULL_MODE_MAX_BYTES = 8 * 1024 * 1024  # larger workbooks are read in streaming mode (no merged-cell info)


def _finish(doc: ParsedDocument, name: str, raw_rows: list[list[Any]], sheet: str | None, extra: dict[str, Any] | None = None) -> bool:
    built = build_table(raw_rows)
    if built is None:
        if any(any(v not in (None, "") for v in r) for r in raw_rows):
            doc.warnings.append(f"{name}: no header row with column names was found; the sheet was not imported")
        return False
    headers, rows, info = built
    if info["skipped_empty_rows"]:
        doc.warnings.append(f"{name}: {info['skipped_empty_rows']} empty row(s) skipped")
    if info["total_rows_skipped"]:
        doc.warnings.append(f"{name}: total row(s) {', '.join(map(str, info['total_rows_skipped']))} skipped")
    doc.add_table(name=name, headers=headers, rows=rows, sheet=sheet, extra={**info, **(extra or {})})
    return True


def parse_csv(data: bytes, filename: str) -> ParsedDocument:
    text, encoding = decode_text(data)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in ",;\t|"}
        delimiter = max(counts, key=counts.get)
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    doc = ParsedDocument(parser="csv", is_tabular=True, meta={"encoding": encoding, "delimiter": delimiter})
    if encoding not in ("utf-8", "ascii"):
        doc.warnings.append(f"File encoding detected as {encoding}")
    _finish(doc, filename, rows, None)
    doc.meta["full_text"] = _render(doc)
    return doc


def parse_xlsx(data: bytes, filename: str) -> ParsedDocument:
    import openpyxl

    streaming = len(data) > FULL_MODE_MAX_BYTES
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=streaming)
    doc = ParsedDocument(parser="openpyxl" + (" (streaming)" if streaming else ""), is_tabular=True,
                         meta={"sheets": wb.sheetnames})
    if streaming:
        doc.warnings.append("Large workbook read in streaming mode: merged cells are not expanded")
    imported = 0
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible":
            doc.warnings.append(f"Hidden sheet '{ws.title}' skipped")
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        merged = 0
        if not streaming and ws.merged_cells.ranges:
            for rng in list(ws.merged_cells.ranges):
                value = ws.cell(rng.min_row, rng.min_col).value
                for r in range(rng.min_row, rng.max_row + 1):
                    for c in range(rng.min_col, rng.max_col + 1):
                        if r - 1 < len(rows):
                            row = rows[r - 1]
                            if c - 1 >= len(row):
                                row.extend([None] * (c - len(row)))
                            row[c - 1] = value
                merged += 1
            doc.warnings.append(f"Sheet '{ws.title}': {merged} merged cell range(s) expanded (value copied to each cell)")
        imported += _finish(doc, f"{filename} / {ws.title}", rows, ws.title, {"merged_ranges": merged})
    wb.close()
    if not imported:
        doc.warnings.append("No importable table was found in the workbook")
    doc.meta["full_text"] = _render(doc)
    return doc


def parse_xls(data: bytes, filename: str) -> ParsedDocument:
    import xlrd

    book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    doc = ParsedDocument(parser="xlrd", is_tabular=True, meta={"sheets": book.sheet_names()})
    for sheet in book.sheets():
        if sheet.visibility:
            doc.warnings.append(f"Hidden sheet '{sheet.name}' skipped")
            continue
        rows: list[list[Any]] = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        for rlo, rhi, clo, chi in sheet.merged_cells:
            value = rows[rlo][clo]
            for r in range(rlo, rhi):
                for c in range(clo, chi):
                    rows[r][c] = value
        if sheet.merged_cells:
            doc.warnings.append(f"Sheet '{sheet.name}': {len(sheet.merged_cells)} merged cell range(s) expanded")
        _finish(doc, f"{filename} / {sheet.name}", rows, sheet.name)
    doc.meta["full_text"] = _render(doc)
    return doc


def _render(doc: ParsedDocument) -> str:
    parts = []
    tables = [s for s in doc.segments if s.kind == "table"]
    for i, t in enumerate(tables):
        seq = doc.segments.index(t)
        rows = [s.content["cells"] for s in doc.segments if s.kind == "row" and s.table_seq == seq][:200]
        parts.append(f"### {t.content['name']}\n" + "\n".join(t.content.get("preamble", [])) + "\n"
                     + markdown_table(t.content["headers"], rows))
        if i > 20:
            break
    return "\n\n".join(parts)
