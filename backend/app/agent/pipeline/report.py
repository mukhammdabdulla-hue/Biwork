"""Stage 9 — Audit & report (docs/architecture.md §6.1): processing summary and XLSX error report."""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.agent.models import AICall, CommittedWrite, ExtractedRecord, IngestionItem, IngestionJob, ParsedSegment, ValidationResult
from app.modules.registry import record_type

SUGGESTIONS = {
    "required_missing": "Fill in the value, or map the spreadsheet column that contains it",
    "invalid_value": "Correct the value to the expected format",
    "reference_not_found": "Choose an existing entry, or add it to master data first",
    "reference_ambiguous": "Choose the intended entry from the candidates",
    "duplicate_record": "Reject this record if it was already recorded; otherwise correct the date/product/unit",
    "duplicate_in_file": "Remove the repeated row or correct it",
    "quantities_exceed_produced": "Check produced / accepted / defective quantities",
    "date_in_future": "Correct the date",
    "conflicting_values": "Confirm which of the conflicting values is correct",
    "unit_incompatible": "Use a unit compatible with the product's base unit",
}


def job_report(session: Session, job: IngestionJob) -> dict[str, Any]:
    records = session.execute(select(ExtractedRecord).where(ExtractedRecord.job_id == job.id).order_by(ExtractedRecord.seq)).scalars().all()
    by_status: dict[str, int] = {}
    groups: dict[str, int] = {}
    fields_extracted = fields_valid = 0
    matched: dict[str, set[str]] = {}
    for r in records:
        by_status[r.status] = by_status.get(r.status, 0) + 1
        if r.status == "needs_review" and r.review_group:
            groups[r.review_group] = groups.get(r.review_group, 0) + 1
        for f in r.fields:
            if f.raw_value not in (None, ""):
                fields_extracted += 1
                if f.state not in ("invalid", "conflicting"):
                    fields_valid += 1
        rt = record_type(r.record_type)
        for name, ref in (r.references or {}).items():
            spec = rt.field(name.split(".")[-1])
            if ref.get("id") and spec is not None:
                matched.setdefault(spec.reference or name, set()).add(ref["id"])
    writes = session.execute(
        select(CommittedWrite, ExtractedRecord.seq).join(ExtractedRecord, ExtractedRecord.id == CommittedWrite.record_id)
        .where(ExtractedRecord.job_id == job.id).order_by(ExtractedRecord.seq)
    ).all()
    ai = session.execute(select(func.count(), func.coalesce(func.sum(AICall.input_tokens), 0),
                                func.coalesce(func.sum(AICall.output_tokens), 0), func.coalesce(func.sum(AICall.cost_usd), 0))
                         .where(AICall.job_id == job.id)).one()
    messages: list[str] = []
    if job.status == "failed":
        messages.append(f"Processing failed at stage '{(job.error or {}).get('stage', job.current_stage)}': "
                        f"{(job.error or {}).get('message', '')}")
    elif records:
        messages.append("Data successfully processed.")
    if records:
        messages.append(f"{len(records)} record(s) extracted.")
        messages.append(f"{fields_extracted} field(s) extracted.")
        messages.append(f"{fields_valid} field(s) validated.")
    for entity, ids in sorted(matched.items()):
        messages.append(f"{len(ids)} existing {entity.replace('_', ' ')}{'' if len(ids) == 1 else 's'} matched.")
    if writes:
        numbers = [w.target_number or str(w.target_id) for w, _ in writes]
        messages.append(f"{len(writes)} record(s) saved: {', '.join(numbers[:10])}{' …' if len(numbers) > 10 else ''}.")
    for status, label in (("approved", "approved, waiting to be saved"), ("failed", "failed to save"), ("rejected", "rejected")):
        if by_status.get(status):
            messages.append(f"{by_status[status]} record(s) {label}.")
    for group, label in (("needs_correction", "need correction"), ("needs_approval", "need approval"), ("ready", "ready to approve")):
        if groups.get(group):
            messages.append(f"{groups[group]} record(s) {label}.")
    return {
        "job_id": str(job.id), "status": job.status, "messages": messages, "records_by_status": by_status, "review_groups": groups,
        "fields_extracted": fields_extracted, "fields_validated": fields_valid,
        "entities_matched": {k: len(v) for k, v in matched.items()},
        "saved": [{"seq": seq, "number": w.target_number, "target_table": w.target_table, "target_id": str(w.target_id),
                   "committed_at": w.committed_at.isoformat()} for w, seq in writes],
        "ai": {"calls": ai[0], "input_tokens": int(ai[1]), "output_tokens": int(ai[2]),
               "cost_usd": float(ai[3]) if isinstance(ai[3], Decimal) else ai[3], "provider": job.ai_provider,
               "model": job.ai_model, "is_mock": job.ai_is_mock},
        "pipeline_version": job.pipeline_version, "prompt_version": job.prompt_version, "warnings": job.warnings,
    }


def error_report_xlsx(session: Session, job: IngestionJob) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Issues"
    headers = ["Record #", "Record type", "Status", "Source", "Field", "Severity", "Problem", "Value", "Suggested fix"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    fills = {"error": PatternFill("solid", fgColor="FDE2E2"), "warning": PatternFill("solid", fgColor="FEF3C7")}
    rows = session.execute(
        select(ValidationResult, ExtractedRecord).join(ExtractedRecord, ExtractedRecord.id == ValidationResult.record_id)
        .where(ExtractedRecord.job_id == job.id, ValidationResult.severity.in_(("error", "warning")),
               ExtractedRecord.status.notin_(("committed", "rejected")))
        .order_by(ExtractedRecord.seq, ValidationResult.id)
    ).all()
    for vr, record in rows:
        ref = record.row_ref or {}
        source = ", ".join(f"{k} {v}" for k, v in ref.items() if k in ("sheet", "row", "page") and v is not None)
        value = None
        if vr.field:
            f = next((x for x in record.fields if x.field == vr.field), None)
            value = f.raw_value if f else None
        ws.append([record.seq, record.record_type, record.status, source, vr.field or "(record)", vr.severity, vr.message,
                   value, SUGGESTIONS.get(vr.rule_code, "")])
        for cell in ws[ws.max_row]:
            cell.fill = fills.get(vr.severity, PatternFill())
    for record in session.execute(select(ExtractedRecord).where(ExtractedRecord.job_id == job.id, ExtractedRecord.status == "failed")).scalars():
        ws.append([record.seq, record.record_type, "failed", "", "(record)", "error", record.error or "Commit failed", None,
                   "Correct the record and approve it again"])
    for item in session.execute(select(IngestionItem).where(IngestionItem.job_id == job.id, IngestionItem.status == "failed")).scalars():
        ws.append(["", "", "failed", item.meta.get("filename") or "text input", "(file)", "error", item.error or "Processing failed",
                   None, "Fix the file and submit it again"])
    for col, width in zip("ABCDEFGHI", (9, 18, 14, 22, 22, 10, 70, 24, 50), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    # Original rows of spreadsheet records with errors, for correction and re-upload.
    tables = session.execute(select(ParsedSegment).join(IngestionItem, IngestionItem.id == ParsedSegment.item_id)
                             .where(IngestionItem.job_id == job.id, ParsedSegment.kind == "table")).scalars().all()
    if tables:
        bad_rows = {(r.row_ref or {}).get("row") for vr, r in rows if vr.severity == "error"}
        sheet = wb.create_sheet("Rows with issues")
        for table in tables:
            sheet.append([f"{table.content['name']}"])
            sheet.append(["Row", *table.content["headers"]])
            for seg in session.execute(select(ParsedSegment).where(ParsedSegment.item_id == table.item_id, ParsedSegment.kind == "row",
                                                                   ParsedSegment.table_seq == table.seq).order_by(ParsedSegment.seq)).scalars():
                if seg.content["row_number"] in bad_rows:
                    sheet.append([seg.content["row_number"], *seg.content["cells"]])
            sheet.append([])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
