"""Builders for test input files."""

from __future__ import annotations

import io
from datetime import date
from typing import Any


def xlsx(rows: list[list[Any]], title: str | None = None, sheet: str = "Hisobot") -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    if title:
        ws.append([title])
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def purchase_request_pdf(number: str, request_date: date, needed_by: date, department: str,
                         items: list[tuple[str, str, str, str]], supplier: str | None = None) -> bytes:
    """A typed (text-layer) purchase request with a ruled item table, in Russian."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
    style = ParagraphStyle("p", fontName="DejaVu", fontSize=11, leading=15)
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4)
    story = [Paragraph(f"ЗАЯВКА НА ЗАКУПКУ № {number}", style), Spacer(1, 6),
             Paragraph(f"Дата: {request_date:%d.%m.%Y}", style),
             Paragraph(f"Подразделение: {department}", style),
             Paragraph(f"Срок поставки: {needed_by:%d.%m.%Y}", style),
             Paragraph("Заявитель: Тошматов Ж.", style)]
    if supplier:
        story.append(Paragraph(f"Поставщик: {supplier}", style))
    story.append(Spacer(1, 10))
    table = Table([["Наименование", "Количество", "Ед. изм", "Цена"], *[list(i) for i in items]])
    table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "DejaVu"), ("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    story += [table, Spacer(1, 10), Paragraph("Обоснование: пополнение запаса на ноябрь", style)]
    doc.build(story)
    return out.getvalue()


def docx_with_table(heading: str, rows: list[list[str]]) -> bytes:
    """A Word document: a heading paragraph and a table (first row = headers)."""
    import docx

    document = docx.Document()
    document.add_paragraph(heading)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def pdf_with_table(heading: str, rows: list[list[str]]) -> bytes:
    """A typed PDF: a heading and a ruled table (first row = headers)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
    out = io.BytesIO()
    table = Table(rows)
    table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "DejaVu"), ("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(out, pagesize=A4).build([Paragraph(heading, ParagraphStyle("p", fontName="DejaVu", fontSize=11)),
                                               Spacer(1, 10), table])
    return out.getvalue()
